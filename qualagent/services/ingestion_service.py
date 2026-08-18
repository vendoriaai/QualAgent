"""Ingestion: parse source files, store originals, and create documents.

Supported formats (04_DATA_MODEL / roadmap 2.1):
- TXT  : plain text
- DOCX : python-docx, paragraphs joined by newlines
- PDF  : pypdf, page text joined by blank lines
- SRT  : srt, timestamps stripped, content joined by newlines
- CSV  : a ``text`` column (required) with optional ``speaker`` column; rows
         with a speaker are emitted as ``Speaker: text`` lines so the ``turns``
         segmenter can recover speakers.

All extracted text is Unicode-NFC normalized with ``\\n`` line endings. The
SHA-256 of the original bytes is stored for provenance.
"""

from __future__ import annotations

import csv
import io
import mimetypes
import unicodedata
from pathlib import Path
from typing import Any

from sqlmodel import Session, col, func, select

from qualagent.domain.errors import DocumentNotFound, ValidationError
from qualagent.domain.models import Document, Segment
from qualagent.services.audit_service import AuditService
from qualagent.services.segmentation_service import SegmentationService, Strategy
from qualagent.storage.files import FileStore

#: Default segmentation strategy applied at import time.
DEFAULT_STRATEGY: Strategy = "sentence"

_EXTENSION_MIME = {
    ".txt": "text/plain",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
    ".srt": "application/x-subrip",
    ".csv": "text/csv",
}


def detect_mime(path: Path) -> str:
    """Return the MIME type for a source path by extension."""
    ext = path.suffix.lower()
    if ext in _EXTENSION_MIME:
        return _EXTENSION_MIME[ext]
    guessed, _ = mimetypes.guess_type(str(path))
    return guessed or "application/octet-stream"


def normalize_text(text: str) -> str:
    """NFC-normalize text and convert line endings to ``\\n``."""
    text = unicodedata.normalize("NFC", text)
    return text.replace("\r\n", "\n").replace("\r", "\n")


def parse_bytes(data: bytes, mime: str, path: Path) -> str:
    """Dispatch to a format parser and return normalized plain text."""
    if mime == "text/plain":
        return normalize_text(_parse_txt(data))
    if mime.endswith("wordprocessingml.document"):
        return normalize_text(_parse_docx(path))
    if mime == "application/pdf":
        return normalize_text(_parse_pdf(path))
    if mime == "application/x-subrip":
        return normalize_text(_parse_srt(data))
    if mime == "text/csv":
        return normalize_text(_parse_csv(data))
    raise ValidationError(f"Unsupported file type for {path.name}: {mime}")


def _parse_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValidationError(f"TXT file is not valid UTF-8: {exc}") from exc


def _parse_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _parse_srt(data: bytes) -> str:
    import srt

    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError(f"SRT file is not valid UTF-8: {exc}") from exc
    subtitles = list(srt.parse(text))
    return "\n".join(sub.content.strip() for sub in subtitles if sub.content.strip())


def _parse_csv(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError(f"CSV file is not valid UTF-8: {exc}") from exc
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None or "text" not in reader.fieldnames:
        raise ValidationError("CSV must have a 'text' column")
    lines: list[str] = []
    for row in reader:
        row_text = (row.get("text") or "").strip()
        if not row_text:
            continue
        speaker = (row.get("speaker") or "").strip()
        lines.append(f"{speaker}: {row_text}" if speaker else row_text)
    return "\n".join(lines)


class IngestionService:
    """Ingest files and manage documents for a project."""

    def __init__(
        self,
        session: Session,
        file_store: FileStore,
        audit: AuditService,
        segmentation: SegmentationService,
    ) -> None:
        self._session = session
        self._files = file_store
        self._audit = audit
        self._segmentation = segmentation

    def ingest_file(
        self,
        project_id: str,
        path: Path,
        *,
        strategy: Strategy = DEFAULT_STRATEGY,
    ) -> tuple[Document, list[Segment]]:
        """Parse, store, persist a document, and run the initial segmentation."""
        data = path.read_bytes()
        sha256 = self._files.store(data)
        mime = detect_mime(path)
        raw_text = parse_bytes(data, mime, path)
        if not raw_text.strip():
            raise ValidationError(f"Document {path.name} has no text content")
        document = Document(
            project_id=project_id,
            filename=path.name,
            sha256=sha256,
            mime=mime,
            raw_text=raw_text,
        )
        self._session.add(document)
        self._session.commit()
        self._session.refresh(document)
        segments = self._segmentation.segment(document, strategy)
        self._audit.emit(
            "system",
            "document.imported",
            {
                "document_id": document.id,
                "filename": document.filename,
                "sha256": sha256,
                "mime": mime,
                "strategy": strategy,
                "segment_count": len(segments),
            },
        )
        return document, segments

    def resegment(
        self,
        document_id: str,
        strategy: Strategy,
    ) -> list[Segment]:
        """Replace a document's segmentation with a new strategy."""
        document = self.get_document(document_id)
        segments = self._segmentation.segment(document, strategy, replace=True)
        self._audit.emit(
            "system",
            "document.resegmented",
            {
                "document_id": document_id,
                "strategy": strategy,
                "segment_count": len(segments),
            },
        )
        return segments

    def get_document(self, document_id: str) -> Document:
        document = self._session.get(Document, document_id)
        if document is None:
            raise DocumentNotFound(f"Document {document_id} not found")
        return document

    def list_documents(
        self,
        project_id: str,
        *,
        limit: int = 50,
        offset: int = 0,
    ) -> tuple[list[Document], int]:
        """List documents for a project with pagination."""
        stmt = select(Document).where(col(Document.project_id) == project_id)
        total = self._session.exec(
            select(func.count()).select_from(Document).where(col(Document.project_id) == project_id)
        ).one()
        stmt = stmt.order_by(col(Document.imported_at)).limit(limit).offset(offset)
        return list(self._session.exec(stmt)), int(total)

    def segment_count(self, document_id: str) -> int:
        """Return the number of segments for a document."""
        return self._session.exec(
            select(func.count()).select_from(Segment).where(col(Segment.document_id) == document_id)
        ).one()

    def to_document_out(self, document: Document) -> dict[str, Any]:
        """Serialize a document with its segment count for API/CLI output."""
        return {
            "id": document.id,
            "project_id": document.project_id,
            "filename": document.filename,
            "sha256": document.sha256,
            "mime": document.mime,
            "imported_at": document.imported_at,
            "segment_count": self.segment_count(document.id),
        }
